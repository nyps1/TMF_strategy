"""
微台10 秒 K 線「依月份分頁」Excel 匯出服務 (services/tick_export/monthly_excel.py)

職責：
1. 以 SQLite (data/market_data.db) 為唯一資料來源 (Single Source of Truth)，
   依「時間」欄位之日曆月份切分，每個月份對應一個工作表 (例如 2026-09 → 工作表「26-09」)。
2. 使用 xlsxwriter constant_memory 串流寫入，記憶體用量與總筆數無關 (僅取決於 chunk_size)。
3. 先寫入暫存檔，成功後以 os.replace 原子替換正本，避免寫入中斷造成正本損毀。
4. 單月筆數若超過 Excel 單表上限 (1,048,576 列)，自動以「26-09 (2)」接續分頁。

設計取捨：xlsx 為 zip 封裝之 XML，無法在不重寫整個封裝的情況下追加列；
本服務採「由 SQLite 全量重建」的冪等策略，換取結構一致與零累積誤差。

遵循 PEP 8、Python Type Hints 與專案工程規範。
"""

import os
import shutil
import sqlite3
import time
import logging
from datetime import datetime
from typing import Any, Dict, Iterator, List, Optional, Tuple

import xlsxwriter

logger = logging.getLogger(__name__)

EXCEL_HEADERS: List[str] = ["時間", "開盤價", "最高價", "最低價", "收盤價", "成交量"]
EXCEL_MAX_DATA_ROWS: int = 1_048_575  # Excel 單表 1,048,576 列，扣除標題列
LEGACY_SHEET_NAME: str = "秒K歷史紀錄"


class MonthlyExcelExporter:
    """
    由 SQLite 秒 K 資料庫產出「一月一工作表」Excel 活頁簿之匯出器。
    """

    def __init__(
        self,
        db_path: str = "data/market_data.db",
        excel_path: str = "data/excel/微台10_每秒逐筆歷史紀錄.xlsx",
        chunk_size: int = 100_000,
        max_rows_per_sheet: int = EXCEL_MAX_DATA_ROWS
    ) -> None:
        """
        初始化月份分頁匯出器。

        參數:
            db_path (str): SQLite 資料庫路徑 (資料來源)。
            excel_path (str): 目標 Excel 活頁簿路徑。
            chunk_size (int): 每次自 SQLite 取回並寫入之列數，控制記憶體峰值。
            max_rows_per_sheet (int): 單一工作表資料列上限 (不含標題列)，預設為 Excel 上限；
                測試時可調小以驗證溢位分頁。
        """
        self.db_path = db_path
        self.excel_path = excel_path
        self.chunk_size = chunk_size
        self.max_rows_per_sheet = max_rows_per_sheet

    # ------------------------------------------------------------------ #
    # 工具方法
    # ------------------------------------------------------------------ #
    @staticmethod
    def month_sheet_name(month_key: str, part: int = 1) -> str:
        """
        將月份鍵轉為工作表名稱。Excel 禁止工作表名稱含「/」，因此以「-」取代 (例如 26/09 → 26-09)。

        參數:
            month_key (str): 月份鍵，格式 YYYY-MM (例如 2026-09)。
            part (int): 分頁序號，1 為主分頁；大於 1 代表溢位接續分頁。

        返回:
            str: 工作表名稱，例如「26-09」或「26-09 (2)」。
        """
        year, month = month_key.split("-")
        base = f"{year[-2:]}-{month}"
        return base if part <= 1 else f"{base} ({part})"

    @staticmethod
    def _month_bounds(month_key: str) -> Tuple[str, str]:
        """
        計算月份之左閉右開時間區間字串。

        參數:
            month_key (str): 月份鍵，格式 YYYY-MM。

        返回:
            Tuple[str, str]: (月初 00:00:00, 次月初 00:00:00)，格式 YYYY-MM-DD HH:MM:SS。
        """
        year, month = (int(x) for x in month_key.split("-"))
        start = datetime(year, month, 1)
        end = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
        return start.strftime("%Y-%m-%d %H:%M:%S"), end.strftime("%Y-%m-%d %H:%M:%S")

    def list_months(self) -> List[Dict[str, Any]]:
        """
        掃描 SQLite 並列出所有含資料之月份及其統計。

        返回:
            List[Dict[str, Any]]: 依月份升冪排列，元素包含 month_key、rows、first_time、last_time。
        """
        conn = sqlite3.connect(self.db_path, timeout=30.0)
        try:
            cur = conn.cursor()
            row = cur.execute("SELECT MIN(time), MAX(time) FROM tmf10_second_bars;").fetchone()
            if not row or not row[0]:
                return []
            first_dt = datetime.strptime(row[0], "%Y-%m-%d %H:%M:%S")
            last_dt = datetime.strptime(row[1], "%Y-%m-%d %H:%M:%S")

            months: List[Dict[str, Any]] = []
            year, month = first_dt.year, first_dt.month
            while (year, month) <= (last_dt.year, last_dt.month):
                key = f"{year:04d}-{month:02d}"
                lo, hi = self._month_bounds(key)
                stats = cur.execute(
                    "SELECT COUNT(*), MIN(time), MAX(time) FROM tmf10_second_bars "
                    "WHERE time >= ? AND time < ?;",
                    (lo, hi)
                ).fetchone()
                if stats and stats[0]:
                    months.append({
                        "month_key": key,
                        "rows": int(stats[0]),
                        "first_time": stats[1],
                        "last_time": stats[2]
                    })
                year, month = (year + 1, 1) if month == 12 else (year, month + 1)
            return months
        finally:
            conn.close()

    def _iter_month_chunks(
        self,
        conn: sqlite3.Connection,
        month_key: str
    ) -> Iterator[List[Tuple[Any, ...]]]:
        """
        以固定大小區塊串流讀取指定月份之秒 K 列 (依時間升冪)。

        參數:
            conn (sqlite3.Connection): 已開啟之 SQLite 連線。
            month_key (str): 月份鍵，格式 YYYY-MM。

        產出:
            List[Tuple[Any, ...]]: 每塊最多 chunk_size 列，欄位順序同 EXCEL_HEADERS。
        """
        lo, hi = self._month_bounds(month_key)
        cur = conn.cursor()
        cur.execute(
            "SELECT time, open, high, low, close, volume FROM tmf10_second_bars "
            "WHERE time >= ? AND time < ? ORDER BY time ASC;",
            (lo, hi)
        )
        while True:
            chunk = cur.fetchmany(self.chunk_size)
            if not chunk:
                break
            yield chunk

    def _backup_legacy_layout(self) -> Optional[str]:
        """
        若現有活頁簿仍為舊版「單一工作表」版型，於首次改版前保留一份 .legacy.bak 備份。

        返回:
            Optional[str]: 備份檔路徑；無需備份或備份已存在時回傳 None。
        """
        if not os.path.exists(self.excel_path):
            return None
        backup_path = f"{self.excel_path}.legacy.bak"
        if os.path.exists(backup_path):
            return None
        try:
            import openpyxl
            wb = openpyxl.load_workbook(self.excel_path, read_only=True)
            names = list(wb.sheetnames)
            wb.close()
        except Exception as e:  # noqa: BLE001
            logger.warning(f"讀取既有活頁簿工作表清單失敗 ({e})，略過舊版備份判斷。")
            return None
        if LEGACY_SHEET_NAME in names:
            shutil.copy2(self.excel_path, backup_path)
            logger.info(f"偵測到舊版單一工作表版型，已備份至: {backup_path}")
            return backup_path
        return None

    # ------------------------------------------------------------------ #
    # 主流程
    # ------------------------------------------------------------------ #
    def export(self) -> Dict[str, Any]:
        """
        由 SQLite 全量重建「一月一工作表」活頁簿，並以原子替換更新正本。

        返回:
            Dict[str, Any]: 成果摘要，包含 sheets (各分頁名稱/筆數/起訖時間)、total_rows、
                file_size_mb、duration_seconds、legacy_backup。

        例外:
            PermissionError: 目標活頁簿正被 Excel 等程式開啟鎖定時，於建置前即快速失敗。
            ValueError: SQLite 內無任何資料。
        """
        t0 = time.time()
        months = self.list_months()
        if not months:
            raise ValueError(
                "Status: ERROR | Root Cause: tmf10_second_bars 無任何資料 | "
                "Suggested Fix: 先執行每日同步或歷史回補後再匯出 Excel。"
            )

        os.makedirs(os.path.dirname(os.path.abspath(self.excel_path)), exist_ok=True)

        # 快速失敗：建置需耗時數分鐘，若正本被鎖定，先行偵測避免白做工。
        if os.path.exists(self.excel_path):
            try:
                with open(self.excel_path, "ab"):
                    pass
            except PermissionError as e:
                raise PermissionError(
                    f"Status: ERROR | Root Cause: 無法寫入 {self.excel_path}，檔案可能正被 Excel 開啟鎖定 ({e}) | "
                    f"Suggested Fix: 關閉正在檢視該檔案之 Excel 後重新執行。"
                ) from e

        legacy_backup = self._backup_legacy_layout()

        tmp_path = f"{self.excel_path}.building.tmp"
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

        workbook = xlsxwriter.Workbook(tmp_path, {"constant_memory": True})
        header_fmt = workbook.add_format({
            "bold": True, "align": "center", "valign": "vcenter",
            "fg_color": "#2c3e50", "font_color": "#ffffff", "border": 1
        })
        time_fmt = workbook.add_format({"align": "center"})
        price_fmt = workbook.add_format({"num_format": "#,##0.0", "align": "right"})
        volume_fmt = workbook.add_format({"num_format": "#,##0", "align": "right"})

        sheet_reports: List[Dict[str, Any]] = []
        total_rows = 0
        last_sheet = None

        conn = sqlite3.connect(self.db_path, timeout=30.0)
        try:
            for m in months:
                part = 1
                ws = None
                row_idx = 0  # 目前工作表已寫入之資料列數
                sheet_rows = 0
                sheet_first: Optional[str] = None
                sheet_last: Optional[str] = None

                def _open_sheet(p: int):
                    """建立月份工作表並寫入標題、欄寬與凍結窗格。"""
                    sheet = workbook.add_worksheet(self.month_sheet_name(m["month_key"], p))
                    for ci, h in enumerate(EXCEL_HEADERS):
                        sheet.write(0, ci, h, header_fmt)
                    sheet.set_column(0, 0, 22, time_fmt)
                    sheet.set_column(1, 4, 12, price_fmt)
                    sheet.set_column(5, 5, 12, volume_fmt)
                    sheet.freeze_panes(1, 0)
                    return sheet

                ws = _open_sheet(part)
                for chunk in self._iter_month_chunks(conn, m["month_key"]):
                    for t, o, h, l, c, v in chunk:
                        if row_idx >= self.max_rows_per_sheet:
                            sheet_reports.append({
                                "sheet": self.month_sheet_name(m["month_key"], part),
                                "rows": sheet_rows, "first_time": sheet_first, "last_time": sheet_last
                            })
                            part += 1
                            ws = _open_sheet(part)
                            row_idx = 0
                            sheet_rows = 0
                            sheet_first = None
                        row_idx += 1
                        ws.write_string(row_idx, 0, t)
                        ws.write_number(row_idx, 1, o)
                        ws.write_number(row_idx, 2, h)
                        ws.write_number(row_idx, 3, l)
                        ws.write_number(row_idx, 4, c)
                        ws.write_number(row_idx, 5, v)
                        sheet_rows += 1
                        if sheet_first is None:
                            sheet_first = t
                        sheet_last = t
                    total_rows += 0  # 於分頁收尾時統一累加
                sheet_reports.append({
                    "sheet": self.month_sheet_name(m["month_key"], part),
                    "rows": sheet_rows, "first_time": sheet_first, "last_time": sheet_last
                })
                last_sheet = ws
                logger.info(
                    f"月份 {m['month_key']} 已寫入工作表「{self.month_sheet_name(m['month_key'], part)}」"
                    f"(本月共 {m['rows']:,} 列，累計 {sum(r['rows'] for r in sheet_reports):,} 列)"
                )
        except Exception:
            conn.close()
            try:
                workbook.close()
            except Exception:  # noqa: BLE001
                pass
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
            raise
        conn.close()

        if last_sheet is not None:
            last_sheet.activate()  # 開啟檔案時預設停在最新月份
        workbook.close()

        total_rows = sum(r["rows"] for r in sheet_reports)
        try:
            os.replace(tmp_path, self.excel_path)
        except PermissionError as e:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
            raise PermissionError(
                f"Status: ERROR | Root Cause: 原子替換 {self.excel_path} 失敗，檔案於建置期間被鎖定 ({e}) | "
                f"Suggested Fix: 關閉正在檢視該檔案之 Excel 後重新執行。"
            ) from e

        size_mb = os.path.getsize(self.excel_path) / (1024.0 * 1024.0)
        duration = time.time() - t0
        logger.info(
            f"月份分頁 Excel 重建完成: {len(sheet_reports)} 個工作表、{total_rows:,} 列、"
            f"{size_mb:.1f} MB、耗時 {duration:.1f} 秒。"
        )
        return {
            "sheets": sheet_reports,
            "total_rows": total_rows,
            "file_size_mb": round(size_mb, 2),
            "duration_seconds": round(duration, 1),
            "legacy_backup": legacy_backup,
            "excel_path": os.path.abspath(self.excel_path)
        }
