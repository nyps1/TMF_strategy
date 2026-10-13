"""
多策略回測成果 Excel 試算表匯出服務 (services/backtest/excel_exporter.py)

負責將多組策略參數之回測結果、逐筆交易明細與時序權益曲線，
匯出為具備金融商務級排版、條件色彩標記與多工作表（Multi-sheet）架構的 .xlsx 檔案。
遵循 PEP 8 與 Python Type Hints 規範。
"""

import os
import logging
from typing import Dict, List, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("ExcelExporter")


class ExcelStrategyExporter:
    """
    提供多策略回測數據整合與高階 Excel 試算表產出服務。
    """

    def __init__(self, filename: str = "three_strategies_trades_comparison.xlsx") -> None:
        """
        初始化 Excel 匯出器。

        參數:
            filename (str): 輸出之 Excel 檔案名稱。
        """
        self.filename = filename
        self.wb = openpyxl.Workbook()
        # 移除預設空白工作表
        self.wb.remove(self.wb.active)

        # 定義商務調色盤
        self.font_family = "Microsoft JhengHei"
        self.header_font = Font(name=self.font_family, size=11, bold=True, color="FFFFFF")
        self.header_fill_navy = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
        self.header_fill_blue = PatternFill(start_color="283593", end_color="283593", fill_type="solid")
        self.header_fill_teal = PatternFill(start_color="004D40", end_color="004D40", fill_type="solid")
        self.header_fill_amber = PatternFill(start_color="E65100", end_color="E65100", fill_type="solid")
        self.sub_header_fill = PatternFill(start_color="263238", end_color="263238", fill_type="solid")
        self.zebra_fill = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")

        self.font_profit = Font(name=self.font_family, size=10, bold=True, color="2E7D32")
        self.font_loss = Font(name=self.font_family, size=10, bold=True, color="C62828")
        self.font_regular = Font(name=self.font_family, size=10)
        self.font_bold = Font(name=self.font_family, size=10, bold=True)

        thin_border_side = Side(border_style="thin", color="E0E0E0")
        self.thin_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side
        )

    def _auto_fit_columns(self, ws) -> None:
        """
        自動適配工作表各欄寬度，確保內容不折行不裁切。

        參數:
            ws: openpyxl Worksheet 物件。
        """
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                # 繁體中文與全形字元依 2 個寬度計算
                line_len = 0
                for ch in val_str:
                    line_len += 2 if ord(ch) > 127 else 1
                if line_len > max_len:
                    max_len = line_len
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    def add_summary_sheet(
        self,
        strategies_summary: List[Dict[str, Any]],
        params_summary: List[Dict[str, Any]]
    ) -> None:
        """
        建立第一分頁：三策略綜合績效對比與參數對照表。

        參數:
            strategies_summary (List[Dict[str, Any]]): 策略量化指標摘要清單。
            params_summary (List[Dict[str, Any]]): 策略參數矩陣清單。
        """
        ws = self.wb.create_sheet(title="綜合績效與參數對照")
        ws.views.sheetView[0].showGridLines = True

        # 1. 標題
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "微台指高頻量化策略 回測績效全方位對比矩陣 (真實成本版)"
        title_cell.font = Font(name=self.font_family, size=16, bold=True, color="1A237E")
        title_cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 36

        # 2. 策略參數矩陣區塊
        ws.cell(row=3, column=1, value="【策略核心參數對照矩陣】").font = Font(name=self.font_family, size=12, bold=True, color="283593")
        param_headers = ["參數項目", "圖一：布林_ATR基準", "圖二：布林_超級趨勢", "圖三：布林_唐奇安", "備註說明"]
        for col_idx, h in enumerate(param_headers, 1):
            c = ws.cell(row=4, column=col_idx, value=h)
            c.font = self.header_font
            c.fill = self.header_fill_navy
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 24

        curr_row = 5
        for p in params_summary:
            ws.cell(row=curr_row, column=1, value=p["item"]).font = self.font_bold
            ws.cell(row=curr_row, column=2, value=p["strat1"]).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=3, value=p["strat2"]).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=4, value=p["strat3"]).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=5, value=p.get("note", "")).font = Font(name=self.font_family, size=9, color="757575")

            for col in range(1, 6):
                cell = ws.cell(row=curr_row, column=col)
                cell.border = self.thin_border
                if curr_row % 2 == 1:
                    cell.fill = self.zebra_fill
            curr_row += 1

        # 3. 核心量化指標對照區塊
        curr_row += 2
        ws.cell(row=curr_row, column=1, value="【核心量化績效指標對照】").font = Font(name=self.font_family, size=12, bold=True, color="004D40")
        curr_row += 1

        kpi_headers = ["量化衡量指標", "圖一：布林_ATR基準", "圖二：布林_超級趨勢", "圖三：布林_唐奇安", "策略表現最佳者", "指標統計意涵"]
        for col_idx, h in enumerate(kpi_headers, 1):
            c = ws.cell(row=curr_row, column=col_idx, value=h)
            c.font = self.header_font
            c.fill = self.header_fill_teal
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[curr_row].height = 24
        curr_row += 1

        for s in strategies_summary:
            ws.cell(row=curr_row, column=1, value=s["metric"]).font = self.font_bold
            c1 = ws.cell(row=curr_row, column=2, value=s["val1"])
            c2 = ws.cell(row=curr_row, column=3, value=s["val2"])
            c3 = ws.cell(row=curr_row, column=4, value=s["val3"])
            c_best = ws.cell(row=curr_row, column=5, value=s["best"])
            c_desc = ws.cell(row=curr_row, column=6, value=s.get("desc", ""))

            c1.alignment = Alignment(horizontal="right")
            c2.alignment = Alignment(horizontal="right")
            c3.alignment = Alignment(horizontal="right")
            c_best.alignment = Alignment(horizontal="center")
            c_best.font = self.font_bold
            c_desc.font = Font(name=self.font_family, size=9, color="757575")

            # 格式化數字
            fmt = s.get("format")
            if fmt:
                c1.number_format = fmt
                c2.number_format = fmt
                c3.number_format = fmt

            for col in range(1, 7):
                cell = ws.cell(row=curr_row, column=col)
                cell.border = self.thin_border
                if curr_row % 2 == 1:
                    cell.fill = self.zebra_fill
            curr_row += 1

        self._auto_fit_columns(ws)

    def add_trades_sheet(
        self,
        sheet_title: str,
        strategy_display_name: str,
        trades: List[Dict[str, Any]],
        theme_fill: PatternFill
    ) -> None:
        """
        建立單一策略之完整逐筆交易明細工作表。

        參數:
            sheet_title (str): 工作表名稱。
            strategy_display_name (str): 策略顯示名稱。
            trades (List[Dict[str, Any]]): 逐筆交易字典清單。
            theme_fill (PatternFill): 表頭主題填充樣式。
        """
        ws = self.wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # 頂部資訊列
        ws.merge_cells("A1:M1")
        title_cell = ws["A1"]
        title_cell.value = f"微台指策略逐筆交易成交稽核表 — {strategy_display_name} (共計 {len(trades)} 筆交易)"
        title_cell.font = Font(name=self.font_family, size=14, bold=True, color="FFFFFF")
        title_cell.fill = theme_fill
        title_cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 32

        # 欄位表頭
        headers = [
            "交易編號", "交易方向", "進場時間", "進場成交價",
            "出場時間", "出場成交價", "在手口數", "出場機制原因",
            "交易摩擦成本(TWD)", "單筆淨損益(TWD)", "單筆報酬率(%)",
            "累計損益(TWD)", "結算後帳戶淨值(TWD)"
        ]

        ws.row_dimensions[2].height = 24
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.sub_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 填入數據
        for row_idx, t in enumerate(trades, 3):
            ws.row_dimensions[row_idx].height = 20
            side_text = "做多 (BUY)" if t["side"] == "LONG" else "放空 (SELL)"
            is_win = t["net_pnl"] >= 0

            c_id = ws.cell(row=row_idx, column=1, value=t["trade_id"])
            c_side = ws.cell(row=row_idx, column=2, value=side_text)
            c_ent_t = ws.cell(row=row_idx, column=3, value=t["entry_time"])
            c_ent_p = ws.cell(row=row_idx, column=4, value=float(t["entry_price"]))
            c_ext_t = ws.cell(row=row_idx, column=5, value=t["exit_time"])
            c_ext_p = ws.cell(row=row_idx, column=6, value=float(t["exit_price"]))
            c_qty = ws.cell(row=row_idx, column=7, value=int(t["qty"]))
            c_reason = ws.cell(row=row_idx, column=8, value=t["exit_reason"])
            c_cost = ws.cell(row=row_idx, column=9, value=float(t["costs"]))
            c_pnl = ws.cell(row=row_idx, column=10, value=float(t["net_pnl"]))
            c_ret = ws.cell(row=row_idx, column=11, value=float(t["return_pct"]) / 100.0)
            c_cum = ws.cell(row=row_idx, column=12, value=float(t["cum_pnl"]))
            c_eq = ws.cell(row=row_idx, column=13, value=float(t["equity_after"]))

            # 對齊與數字格式
            c_id.alignment = Alignment(horizontal="center")
            c_side.alignment = Alignment(horizontal="center")
            c_ent_t.alignment = Alignment(horizontal="center")
            c_ent_p.alignment = Alignment(horizontal="right")
            c_ent_p.number_format = "#,##0.0"
            c_ext_t.alignment = Alignment(horizontal="center")
            c_ext_p.alignment = Alignment(horizontal="right")
            c_ext_p.number_format = "#,##0.0"
            c_qty.alignment = Alignment(horizontal="center")
            c_reason.alignment = Alignment(horizontal="center")
            c_cost.alignment = Alignment(horizontal="right")
            c_cost.number_format = "#,##0"

            c_pnl.alignment = Alignment(horizontal="right")
            c_pnl.number_format = "$#,##0;($#,##0);\"-\""
            c_pnl.font = self.font_profit if is_win else self.font_loss

            c_ret.alignment = Alignment(horizontal="right")
            c_ret.number_format = "+0.00%;-0.00%;0.00%"
            c_ret.font = self.font_profit if is_win else self.font_loss

            c_cum.alignment = Alignment(horizontal="right")
            c_cum.number_format = "#,##0.0"
            c_cum.font = self.font_profit if t["cum_pnl"] >= 0 else self.font_loss

            c_eq.alignment = Alignment(horizontal="right")
            c_eq.number_format = "#,##0.0"
            c_eq.font = self.font_bold

            # 邊框與斑馬紋
            for col in range(1, 14):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = self.thin_border
                if row_idx % 2 == 1:
                    cell.fill = self.zebra_fill

        self._auto_fit_columns(ws)

    def add_equity_curve_sheet(
        self,
        time_series: List[Dict[str, Any]]
    ) -> None:
        """
        建立時序權益曲線對照工作表。

        參數:
            time_series (List[Dict[str, Any]]): 包含 datetime 與三策略權益之字典清單。
        """
        ws = self.wb.create_sheet(title="時序淨值對照")
        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("A1:E1")
        title_cell = ws["A1"]
        title_cell.value = "三策略時間序列累積損益與淨值推進對照表"
        title_cell.font = Font(name=self.font_family, size=14, bold=True, color="FFFFFF")
        title_cell.fill = self.header_fill_navy
        title_cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 32

        headers = ["時間 (Datetime)", "收盤價 (Close)", "布林_ATR基準 (TWD)", "布林_超級趨勢 (TWD)", "布林_唐奇安 (TWD)"]
        ws.row_dimensions[2].height = 24
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.sub_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, pt in enumerate(time_series, 3):
            c_time = ws.cell(row=row_idx, column=1, value=pt["datetime"])
            c_close = ws.cell(row=row_idx, column=2, value=float(pt["close"]))
            c_s1 = ws.cell(row=row_idx, column=3, value=float(pt["strat1_equity"]))
            c_s2 = ws.cell(row=row_idx, column=4, value=float(pt["strat2_equity"]))
            c_s3 = ws.cell(row=row_idx, column=5, value=float(pt["strat3_equity"]))

            c_time.alignment = Alignment(horizontal="center")
            c_close.alignment = Alignment(horizontal="right")
            c_close.number_format = "#,##0.0"

            for c in [c_s1, c_s2, c_s3]:
                c.alignment = Alignment(horizontal="right")
                c.number_format = "#,##0.0"

            for col in range(1, 6):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = self.thin_border
                if row_idx % 2 == 1:
                    cell.fill = self.zebra_fill

        self._auto_fit_columns(ws)

    def save(self) -> str:
        """
        儲存試算表至本機並回傳絕對路徑。
        具備 Windows 檔案鎖定容錯重試機制，若檔案被外部 Excel 軟體佔用則自動切換至備用檔名。

        返回:
            str: 產出之 Excel 檔案絕對路徑。
        """
        import time
        abs_path = os.path.abspath(self.filename)
        for attempt in range(3):
            try:
                self.wb.save(abs_path)
                logger.info(f"Excel 試算表成功匯出: {abs_path}")
                return abs_path
            except PermissionError:
                if attempt < 2:
                    time.sleep(0.5)
                    continue
                # 第三次仍被佔用，儲存為獨立時間戳備用檔名避免管線中斷
                backup_name = abs_path.replace(".xlsx", f"_backup_{int(time.time())}.xlsx")
                self.wb.save(backup_name)
                logger.warning(f"檔案 {abs_path} 正被外部程式佔用鎖定，已自動儲存至替代路徑: {backup_name}")
                return backup_name
